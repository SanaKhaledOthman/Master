#!/usr/bin/env python3
"""Build the systematic-review Excel workbook from a review JSON file.

Usage:
    python build_workbook.py review.json review.xlsx

Requires openpyxl (pip install openpyxl). See SKILL.md, section 7, for the
JSON layout. Missing keys are written as "Not reported".
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NR = "Not reported"

# (JSON key, column header, column width)
EXTRACTION_COLUMNS = [
    ("id", "ID", 6),
    ("title", "Title", 40),
    ("authors", "Authors", 25),
    ("year", "Year", 7),
    ("venue", "Journal", 28),
    ("publisher", "Publisher", 14),
    ("quartile", "Quartile", 9),
    ("sjr", "SJR", 8),
    ("impact_factor", "Impact factor", 10),
    ("ranking_source", "Ranking source", 20),
    ("doi_url", "DOI / URL", 30),
    ("problem", "Problem", 40),
    ("objective", "Objective", 40),
    ("method", "Method", 22),
    ("model", "Model/Algorithm", 28),
    ("dataset", "Dataset", 30),
    ("features", "Features", 30),
    ("labels", "Labels", 22),
    ("threat_type", "Attack/Threat Type", 28),
    ("evaluation", "Evaluation", 30),
    ("results", "Results", 35),
    ("baseline", "Baseline", 30),
    ("limitations", "Limitations", 40),
    ("future_work", "Future Work", 35),
    ("generalization", "Generalization", 30),
    ("real_world", "Real-world applicability", 30),
    ("explainability", "Explainability", 25),
    ("reproducibility", "Reproducibility", 25),
    ("source_note", "Extraction source", 16),
    ("quality_total", "Quality (/14)", 10),
]

SUMMARY_COLUMNS = [
    ("id", "ID", 6), ("title", "Title", 40), ("year", "Year", 7),
    ("venue", "Journal", 28), ("quartile", "Quartile", 9),
    ("method", "Method", 22), ("model", "Model", 25), ("dataset", "Dataset", 25),
    ("labels", "Labels", 18), ("threat_type", "Threat", 22),
    ("best_result", "Best result", 30), ("generalization", "Generalization", 22),
    ("explainability", "XAI", 14), ("reproducibility", "Code/Data", 16),
]

QUALITY_CRITERIA = [
    ("clear_problem", "Clear problem"),
    ("sound_split", "Leakage-free split"),
    ("strong_baselines", "Strong baselines"),
    ("imbalance_metrics", "Imbalance-aware metrics"),
    ("multiple_datasets", "Multiple datasets"),
    ("statistics", "Seeds / variance"),
    ("code_data", "Code / data available"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def cell_value(value):
    if value is None or value == "":
        return NR
    if isinstance(value, (list, tuple)):
        return "; ".join(str(v) for v in value) or NR
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items()) or NR
    return value


def write_table(wb, name, columns, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    for col, (_, header, width) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, WRAP
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=2):
        for col, (key, _, _) in enumerate(columns, start=1):
            c = ws.cell(row=r, column=col, value=cell_value(row.get(key)))
            c.alignment = WRAP
    ws.freeze_panes = "C2" if len(columns) > 4 else "A2"
    if rows:
        ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        table = Table(displayName=name.replace(" ", "_").replace("-", "_"), ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    return ws


def main(src, dst):
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    papers = data.get("papers", [])
    for i, p in enumerate(papers, start=1):
        p.setdefault("id", i)
        q = p.get("quality") or {}
        if "quality_total" not in p and q:
            p["quality_total"] = sum(int(v) for v in q.values() if str(v).isdigit())

    wb = Workbook()

    # Protocol
    proto = data.get("protocol", {})
    rows = [
        {"field": "Topic", "value": data.get("topic")},
        {"field": "Search date", "value": data.get("search_date")},
        {"field": "Research question", "value": proto.get("question")},
        {"field": "Journal ranking rule", "value": proto.get("ranking_rule")},
        {"field": "Inclusion criteria", "value": proto.get("inclusion")},
        {"field": "Exclusion criteria", "value": proto.get("exclusion")},
        {"field": "Tools used", "value": proto.get("tools_used")},
        {"field": "Tools unavailable", "value": proto.get("tools_unavailable")},
    ]
    write_table(wb, "Protocol", [("field", "Field", 24), ("value", "Value", 100)], rows, first=True)
    write_table(wb, "Search Log",
                [("tool", "Tool", 22), ("query", "Query", 70), ("hits", "Hits", 8), ("date", "Date", 12)],
                proto.get("queries", []))
    write_table(wb, "PRISMA",
                [("stage", "Stage", 50), ("count", "Count", 10), ("note", "Note", 60)],
                data.get("prisma", []))
    write_table(wb, "Extraction", EXTRACTION_COLUMNS, papers)
    write_table(wb, "Summary", SUMMARY_COLUMNS, papers)

    qcols = [("id", "ID", 6), ("title", "Title", 40)] + \
        [(k, h, 12) for k, h in QUALITY_CRITERIA] + [("quality_total", "Total (/14)", 10)]
    qrows = [{"id": p["id"], "title": p.get("title"), "quality_total": p.get("quality_total"),
              **(p.get("quality") or {})} for p in papers]
    write_table(wb, "Quality", qcols, qrows)

    write_table(wb, "Gap Matrix",
                [("axis", "Gap axis", 22), ("covered", "Covered", 40), ("missing", "Missing", 40),
                 ("evidence", "Evidence (counts)", 35), ("papers", "Papers", 20)],
                data.get("gaps", []))
    write_table(wb, "Opportunities",
                [("direction", "Research direction", 60), ("gaps", "Addresses gaps", 30),
                 ("papers", "Motivating papers", 25)],
                data.get("opportunities", []))
    write_table(wb, "Excluded",
                [("title", "Title", 50), ("venue", "Venue", 30), ("quartile", "Quartile", 9),
                 ("stage", "Stage", 18), ("reason", "Reason", 50)],
                data.get("excluded", []))
    refs = [{"id": p["id"], "citation": f'{p.get("authors", NR)} ({p.get("year", "n.d.")}). '
             f'{p.get("title", NR)}. {p.get("venue", NR)}.', "url": p.get("doi_url")} for p in papers]
    write_table(wb, "References",
                [("id", "ID", 6), ("citation", "Citation", 90), ("url", "DOI / URL", 40)], refs)

    wb.save(dst)
    print(f"Wrote {dst}: {len(papers)} included, {len(data.get('excluded', []))} excluded")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
