"""Adds two sheets to systematic_review.xlsx:

* "All Papers - Gap Test": every included, excluded and later-found paper (preprints too), with the
  components it covers, which part of the planned work it relates to, and a duplication test
  against each of the four candidate gaps.
* "Gap Duplication Summary": per-gap verdict, closest papers and what remains novel.

Run after build_workbook.py:  python add_gap_test.py
"""
import json
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "systematic_review.xlsx")

GAPS = {
    "G1": "XAI-gated pseudo-labels: explanations decide whether an automatically labelled zero-day sample may enter continual retraining",
    "G2": "Explanation-aware DRL controller: a DRL agent uses explanation signals (state or reward) to decide when / how much to retrain a detector, with veto or rollback",
    "G3": "Explaining continual updates: explanations that show how the detector changes across updates (forgetting, drift impact)",
    "G4": "Encoding vs novelty: how categorical feature encoding affects unsupervised zero-day detectability",
}

NO = ("No overlap", "")
DUP, HIGH, PART, LOW = "Duplicates gap", "High overlap", "Partial overlap", "Low overlap"

# (title, year, venue, venue_type, U, D, C, X, related_part, g1, g2, g3, g4, extraction_source, link)
P = []


def add(title, year, venue, vtype, comps, part, g1=NO, g2=NO, g3=NO, g4=NO, src="abstract", link=""):
    P.append(dict(title=title, year=year, venue=venue, vtype=vtype,
                  U="✓" if "U" in comps else "", D="✓" if "D" in comps else "",
                  C="✓" if "C" in comps else "", X="✓" if "X" in comps else "",
                  part=part, g1=g1, g2=g2, g3=g3, g4=g4, src=src, link=link))


Q1 = "Q1 journal (included)"
# ---------------- included (Q1 journals) ----------------
add("GARDIAN: Deep RL for autonomous and continual network intrusion detection", 2026, "J. Information Security and Applications", Q1, "UDC",
    "Baseline framework for your work: unsupervised labelling + DRL + continual learning (no XAI)",
    g1=(PART, "Pseudo-labels from MAE + DBSCAN with no quality gate and no XAI; our run: 82.5% correct, harmful for fine-tuning"),
    g2=(PART, "PPO decides when / how much to retrain, but its state has no explanation signals and it cannot roll back"),
    src="full text", link="https://doi.org/10.1016/j.jisa.2026.104594")
add("A Self-Adaptive IDS for Zero-Day Attacks Using Deep Q-Networks", 2025, "IEEE Access", Q1, "D",
    "Zero-day ransomware DRL baseline on UGRansome",
    g4=(PART, "Uses label-encoded categoricals without studying encoding; our test shows it hides unseen families (AUC 0.49)"),
    src="full text", link="https://doi.org/10.1109/ACCESS.2025.3617792")
add("An active learning framework using deep Q-network for zero-day attack detection", 2024, "Computers & Security", Q1, "D",
    "DRL chooses which zero-day samples get labelled",
    g1=(PART, "DQN selects samples for labelling; labels assigned by distance, never checked with explanations"),
    g2=(LOW, "DRL used for sample selection, not to control retraining"), link="https://doi.org/10.1016/j.cose.2024.103713")
add("Auto-IDaS: RL to configure models, tasks and capacities", 2024, "J. Network and Computer Applications", Q1, "D",
    "DRL as an IDS manager (model selection)", g2=(LOW, "DQN switches detector models over time; no XAI, no retraining decision"),
    link="https://doi.org/10.1016/j.jnca.2024.103936")
add("RL for Intrusion Detection: More Model Longness and Fewer Updates", 2023, "IEEE Trans. Network and Service Management", Q1, "DC",
    "Long-horizon RL classifier with transfer-learning updates (4 years of real traffic)",
    g2=(LOW, "Update schedule is periodic; RL does not decide when to update; no XAI"), src="full text (partial)",
    link="https://secplab.ppgia.pucpr.br/files/papers/2023tnsm.pdf")
add("Adaptive Defense: Zero-Day Attack Detection in NIDS with DRL", 2025, "IEEE Access", Q1, "D",
    "DRL zero-day baseline with public code", link="https://ieeexplore.ieee.org/document/11063272")
add("DQ-IDS: Deep Q-learning intrusion detection system", 2025, "ICT Express", Q1, "D", "DRL classifier baseline",
    link="https://www.sciencedirect.com/science/article/pii/S2405959525000694")
add("Cross-dataset zero-day IDS integrating Siamese network and RL", 2026, "ICT Express", Q1, "UD",
    "Unsupervised detection + PPO; lists XAI as future work",
    g1=(PART, "Zero-day labels defined by a distance threshold, never validated"), src="full text (arXiv version)",
    link="https://www.sciencedirect.com/science/article/pii/S2405959526000731")
add("ULTIMATE: Multi-agent DRL for false-positive optimized enterprise ID", 2026, "Array", Q1, "DX",
    "DRL + SHAP example", g2=(LOW, "SHAP explains the agents after the fact; explanations do not drive decisions"),
    link="https://www.sciencedirect.com/science/article/pii/S2590005626002195")
add("Adaptive IIoT intrusion detection using DRL against APT attacks (DRL-GCN)", 2026, "Eng. Applications of AI", Q1, "DC",
    "Online DRL policy adaptation under drift", g2=(LOW, "DRL updates its policy online; no explanation signals"),
    link="https://www.sciencedirect.com/science/article/abs/pii/S0952197626018336")
add("IGPC-MSOS: knowledge-preserving mode-switching for concept drift", 2026, "Knowledge-Based Systems", Q1, "C",
    "Drift-triggered adaptation baseline", g2=(LOW, "Drift + performance trigger decides adaptation (not RL, not XAI)"),
    link="https://doi.org/10.1016/j.knosys.2026.115361")
add("RL-based voting for feature drift-aware ID (IFDA-GPC)", 2025, "IEEE Access", Q1, "DC",
    "DQN feature selection under feature drift", g2=(LOW, "DRL picks features, not retraining; no XAI"),
    link="https://ieeexplore.ieee.org/document/10896652")
add("Analysis of Continual Learning Models for IDS", 2022, "IEEE Access", Q1, "C",
    "Continual-learning baselines (LwF, ER, DER)", g3=(LOW, "Measures forgetting but gives no explanations"),
    link="https://doi.org/10.1109/ACCESS.2022.3222715")
add("RepShield: robust knowledge representation in CL for NIDS", 2026, "Computer Networks", Q1, "C",
    "Continual-learning baseline with drift detection", g3=(LOW, "PCA drift detector; no explanations"),
    link="https://www.sciencedirect.com/science/article/abs/pii/S1389128626003208")
add("Latent space alignment for IoT botnet detection in non-stationary environments", 2025, "Knowledge-Based Systems", Q1, "C",
    "Drift robustness without retraining", link="https://www.sciencedirect.com/science/article/pii/S0950705125017873")
add("VLSA-CL: variational latent alignment + contrastive learning under drift", 2026, "Neurocomputing", Q1, "C",
    "Drift robustness without retraining", link="https://www.sciencedirect.com/science/article/pii/S0925231226014967")
add("ONIDS: class-imbalance and drift invariant online botnet detection", 2024, "Computers & Security", Q1, "C",
    "Online drift-adaptive IoT IDS", link="https://www.sciencedirect.com/science/article/abs/pii/S0167404824001214")
add("An adaptable deep learning-based IDS to zero-day attacks (DOC++)", 2023, "J. Information Security and Applications", Q1, "UC",
    "Open-set detection, then clustering of unknown samples for later updates",
    g1=(PART, "Clusters novel samples so they can be labelled; no explanation-based validation"),
    link="https://doi.org/10.1016/j.jisa.2023.103516")
add("Intelligent zero-day attack detection using unsupervised ML (ZdAD-UML)", 2025, "Knowledge-Based Systems", Q1, "UX",
    "Unsupervised zero-day + SHAP", g1=(LOW, "SHAP used to explain detections only; no labels, no updates"),
    link="https://doi.org/10.1016/j.knosys.2025.113833")
add("Semi-supervised in-vehicle IDS with VAE and adversarial RL", 2024, "Knowledge-Based Systems", Q1, "UD",
    "Pseudo-labels + RL for unknown attacks", g1=(PART, "Pseudo-labels for pre-training without an explanation check"),
    g2=(LOW, "Adversarial RL agents classify / select samples, not retraining control"), link="https://doi.org/10.1016/j.knosys.2024.112563")
add("Application of DRL to intrusion detection for supervised problems", 2020, "Expert Systems with Applications", Q1, "D",
    "Foundational DRL-IDS reference")
add("CNN feature extraction + OC-SVM for zero-day detection", 2026, "Scientific Reports", Q1, "U",
    "One-class zero-day baseline with code", src="full text (PMC)", link="https://pmc.ncbi.nlm.nih.gov/articles/PMC13527146/")

# ---------------- preprints ----------------
PRE = "Preprint (arXiv)"
add("Enhancing Autonomous Online IDS for IoT with Reliable Pseudo-Labels (PseudoFilter)", 2026, "arXiv 2605.26166", PRE, "UC",
    "Closest non-XAI pseudo-label gate: your G1 baseline",
    g1=(HIGH, "Accepts a pseudo-label only if confidence >= 0.85 AND encoder/decoder agree (+1.25 pts acc); uses no XAI. Also shows AOC-IDS accuracy collapsing 90% -> 75% without a gate"),
    src="full text", link="https://arxiv.org/abs/2605.26166")
add("SOUL: Semi-supervised Open-world continUal Learning for NIDS", 2024, "arXiv 2412.00911", PRE, "UC",
    "Label-efficient CL baseline (<=20% labels)",
    g1=(HIGH, "Generates high-confidence labels for new tasks; confidence-based, no XAI"), link="https://arxiv.org/abs/2412.00911")
add("RSST-NIDS: Robust Semi-Supervised Temporal ID for adversarial cloud networks", 2026, "arXiv 2604.12655", PRE, "UC",
    "Confidence-aware pseudo-labelling under drift", g1=(HIGH, "Confidence-aware pseudo-labelling; no XAI"),
    link="https://arxiv.org/abs/2604.12655")
add("UNAD+: Explainable Hybrid Framework for Unknown Network Attack Detection", 2026, "arXiv 2605.22621", PRE, "UX",
    "Unsupervised zero-day -> pseudo-labels -> supervised refinement + LIME",
    g1=(HIGH, "Pseudo-labels refine a classifier; a HUMAN checks a subset; LIME/surrogate are post hoc and never gate labels"),
    src="full text", link="https://arxiv.org/abs/2605.22621")
add("ACORN-IDS: Adaptive Continual Novelty Detection for IDS", 2026, "arXiv 2602.07291", PRE, "UC",
    "Continual novelty detection; strong unsupervised-CL baseline (5 datasets)",
    g1=(PART, "K-Means pseudo-labels from overlap with clean normal data; no gate, no XAI"),
    g4=(LOW, "One-hot encodes categoricals but does not study the effect"), src="full text", link="https://arxiv.org/abs/2602.07291")
add("CITADEL: Continual Anomaly Detection for IoT Intrusion Detection", 2025, "arXiv 2508.19450", PRE, "UC",
    "Self-supervised CL novelty-detection baseline", link="https://arxiv.org/abs/2508.19450")
add("ThreatFormer-IDS: zero-day generalization and explainable attribution", 2026, "arXiv 2603.00185", PRE, "UX",
    "Zero-day + Integrated-Gradients attribution per alert",
    g3=(LOW, "Explains individual alerts; CL listed as future work"), g4=(LOW, "Categorical vocabulary embeddings, not studied"),
    src="full text", link="https://arxiv.org/abs/2603.00185")
add("DeepXplain: XAI-Guided Autonomous Defense Against Multi-Stage APT", 2026, "arXiv 2603.21296 (under review, GLOBECOM 2026)", PRE, "DX",
    "Explanation signals inside DRL training: closest to G2",
    g2=(HIGH, "Explanation-alignment loss + explanation-confidence reward shape a PPO defence policy. BUT it chooses response actions for APTs, not detector retraining; no CL, no rollback"),
    src="full text", link="https://arxiv.org/abs/2603.21296")
add("C-MADF: Explainable Autonomous Cyber Defense using Adversarial MARL", 2026, "arXiv 2604.04442", PRE, "DX",
    "Explainability score gates autonomous actions",
    g2=(HIGH, "Explainability-Transparency Score decides escalation vs autonomous action; for response actions, not retraining"),
    src="full text", link="https://arxiv.org/abs/2604.04442")
add("Navigating the Latent Manifold: Proactive Concept Drift Adaptation for NIDS", 2026, "arXiv 2609.08623", PRE, "C",
    "Drift adaptation related work", link="https://arxiv.org/abs/2609.08623")
add("Universal adversarial perturbations with explainability against DRL-based IDS", 2026, "arXiv 2609.30605", PRE, "DX",
    "Robustness of DRL-IDS", g2=(LOW, "Explains attacks on a DRL-IDS; not adaptation"), link="https://arxiv.org/abs/2609.30605")
add("HuntGPT: ML anomaly detection + XAI + LLM", 2023, "arXiv 2309.16021", PRE, "X", "Analyst-facing XAI",
    link="https://arxiv.org/abs/2309.16021")
add("Explain to Not Forget: Defending Against Catastrophic Forgetting with XAI", 2022, "arXiv 2205.01929 (venue not verified)", PRE, "CX",
    "XAI against forgetting (outside IDS)",
    g3=(HIGH, "Uses explanations to reduce forgetting in CL, but on vision tasks, not NIDS"), link="https://arxiv.org/abs/2205.01929")
add("CIP-Net: Continual Interpretable Prototype-based Network", 2025, "arXiv 2512.07981", PRE, "CX",
    "Interpretable continual learning (outside IDS)", g3=(PART, "Interpretable prototypes in CL; not security"),
    link="https://arxiv.org/abs/2512.07981")
add("ADAPT: pseudo-labeling to combat concept drift in malware detection", 2025, "arXiv 2507.08597", PRE, "UC",
    "Pseudo-labels for drift (malware, not NIDS)", g1=(PART, "Pseudo-labelling for drift adaptation; no XAI gate in title/abstract"),
    src="title/abstract snippet", link="https://arxiv.org/abs/2507.08597")

# ---------------- conferences / workshops / chapters ----------------
CONF = "Conference / workshop / chapter"
add("AOC-IDS: Autonomous Online Framework with Contrastive Learning", 2024, "IEEE INFOCOM 2024", CONF, "UC",
    "Autonomous online IDS with AE pseudo-labels; code available",
    g1=(PART, "Pseudo-labels with no quality gate (only 5% random flips); accuracy collapses over time per Afzaal et al."),
    link="https://github.com/xinchen930/AOC-IDS")
add("SSF: Continual Learning with Strategic Selection and Forgetting for NIDS", 2025, "IEEE INFOCOM 2025", CONF, "C",
    "Strong CL baseline (top-tier venue)", link="https://arxiv.org/abs/2412.16264")
add("Online IDS framework with explainable concept drift detection and adaptation (Drago et al.)", 2026, "ITASEC & SERICS 2026 (CEUR-WS 4198)", CONF, "UCX",
    "Closest to G3; partly G2 (XAI decides when to adapt)",
    g1=(LOW, "Lowers labelling cost by requesting few human labels; no pseudo-label validation"),
    g2=(PART, "Label-free drift detector based on changes in SHAP feature-importance profiles triggers adaptation; no RL"),
    g3=(HIGH, "Explains the drift's impact and shows how model decisions evolve over time"),
    src="full text", link="https://ceur-ws.org/Vol-4198/paper54.pdf")
add("ExpIDS: drift-adaptable NIDS with improved explainability", 2024, "IEEE conference (Xplore 10815461)", CONF, "UCX",
    "Drift-adaptive AE with decision-tree explanations",
    g3=(PART, "Global TRUSTEE explanations of the current model; drift updates via U-index with ground-truth labels; does not explain changes between updates"),
    src="full text", link="https://arxiv.org/abs/2509.20767")
add("CADE: Detecting and Explaining Concept Drift Samples for Security Applications", 2021, "USENIX Security 2021", CONF, "UCX",
    "Explains drifting samples for analysts",
    g1=(PART, "Explains why a sample drifted to help analysts label it; no automatic gate"), g3=(HIGH, "Detects and explains concept-drift samples"),
    src="title + citation (ExpIDS ref. 9)", link="https://www.usenix.org/conference/usenixsecurity21/presentation/yang-limin")
add("ENIDrift: fast adaptive ensemble NIDS under real-world drift", 2022, "ACSAC 2022", CONF, "C", "Drift-adaptive NIDS baseline",
    src="citation (ExpIDS ref. 13)", link="https://doi.org/10.1145/3564625.3567992")
add("Replay or Regret: Evaluating CL Methods for Robust ID", 2025, "IEEE MILCOM 2025", CONF, "C", "Replay is the most effective CL strategy",
    g3=(LOW, "Measures forgetting, no explanations"))
add("SCL-IDS: Semi-Supervised Continual Learning for Adaptive ID (poster)", 2025, "IEEE ICNP 2025", CONF, "UC", "Semi-supervised CL",
    g1=(LOW, "Semi-supervised CL with scarce labels; no XAI"))
add("MAML AE-SAC: Meta-Learning and Adversarial RL for Adaptive ID", 2025, "KSE 2025", CONF, "D", "Fast-adapting RL IDS",
    g2=(LOW, "RL adapts quickly; no explanation signals"))
add("Q-ID: RL framework for adaptive intrusion detection", 2025, "FedCSIS 2025", CONF, "D", "RL classifier")
add("DRL-based IDS for next-generation wireless networks", 2025, "ICIRCA 2025", CONF, "D", "DRL classifier")
add("Network anomaly detection using Q-Learning and DQN (UNSW-NB15 zero-day)", 2025, "ICSECS 2025", CONF, "D", "DRL zero-day classifier")
add("ZDADS: zero-day detection in cloud with adaptive DQN", 2025, "ICCES 2025", CONF, "D", "DRL zero-day classifier")
add("Explainable graph-based RL for intrusion detection", 2025, "ICCMC 2025", CONF, "DX", "Conceptual XAI + RL",
    g2=(LOW, "Proposes combining XAI with RL conceptually; no explanation-driven decisions"))
add("RL for automated ID and adaptive defense in zero-day scenarios (chapter)", 2025, "Book chapter", CONF, "D", "Background")

# ---------------- journals below Q1 ----------------
LOWQ = "Journal below Q1 / rank not verified"
add("Continual Learning for Intrusion Detection Under Evolving Network Threats (Guo et al.)", 2025, "Future Internet (Q2)", LOWQ, "UC",
    "Clustering memory + distillation + pseudo-label filtering",
    g1=(HIGH, "Confidence-aware pseudo-label filtering (entropy + agreement); no XAI"), link="https://doi.org/10.3390/fi17100456")
add("Post-hoc categorization based on XAI and RL for improved ID", 2024, "Applied Sciences (Q2)", LOWQ, "DX",
    "Explaining an RL-based IDS", g2=(LOW, "Multi-level post-hoc explanations of an RL classifier"),
    link="https://doi.org/10.3390/app142411511")
add("Continual learning process for learned and emerging attacks (Memory-LGBM)", 2025, "Applied Sciences (Q2)", LOWQ, "C",
    "Prototype-memory CL baseline", link="https://doi.org/10.3390/app151810034")
add("Versatile XAI-based framework for efficient explainable IDS", 2025, "Annals of Telecommunications (Q2)", LOWQ, "X",
    "Cost and reliability of XAI (SHAP vs LIME cross-check)",
    g1=(LOW, "Cross-checking two explainers is a reliability test that could serve inside a gate"),
    link="https://doi.org/10.1007/s12243-025-01118-9")
add("SpIDER: satellite ID using explainable RL", 2026, "Telecom (MDPI)", LOWQ, "DX", "DRL + SHAP example",
    g2=(LOW, "SHAP explains the RL model post hoc"), link="https://doi.org/10.3390/telecom7010003")
add("Harnessing DBSCAN and autoencoder for intrusion detection in cloud", 2024, "Bull. Electrical Eng. & Informatics", LOWQ, "U",
    "AE detection + DBSCAN clustering of attacks", g1=(LOW, "Clusters attacks; no labels for retraining, no XAI"),
    link="https://doi.org/10.11591/eei.v13i5.8135")
add("XAI-based model drift detection applicable to unsupervised environments", 2023, "Computers, Materials & Continua (rank not verified)", LOWQ, "UCX",
    "SHAP-based drift detection without labels",
    g2=(PART, "Explanations decide when drift occurred; no RL"), g3=(HIGH, "Uses SHAP to detect and characterise drift"),
    link="https://www.sciencedirect.com/org/science/article/pii/S1546221823001558")

# ---------------- reviews / out of domain ----------------
REV = "Review / out of domain"
add("Application of DRL for intrusion detection in IoT: a systematic review", 2025, "Internet of Things (Q1)", REV, "D", "Background on DRL-IDS")
add("Evolving cybersecurity frontiers: survey on concept drift in IDS", 2024, "Eng. Applications of AI (Q1)", REV, "C", "Background on drift")
add("XAI-Guided Continual Learning: Rationale, Methods, and Future Directions", 2025, "WIREs Data Mining & Knowledge Discovery", REV, "CX",
    "Shows XAI-for-CL exists outside security", g3=(HIGH, "Reviews XAI used against forgetting (general ML, not IDS)"),
    link="https://wires.onlinelibrary.wiley.com/doi/10.1002/widm.70046")
add("Explaining anomalies detected by autoencoders using SHAP (Antwarg et al.)", 2021, "Expert Systems with Applications (Q1)", REV, "UX",
    "Method you would build G1 on", g1=(PART, "Explains AE anomalies with kernel SHAP; no labels or retraining (enabler, not duplicate)"),
    link="https://www.sciencedirect.com/science/article/abs/pii/S0957417421011155")

SUMMARY = [
    ("G1", GAPS["G1"], "OPEN",
     "Afzaal et al. 2026 (confidence + agreement gate); Guo et al. 2025 (entropy + agreement); SOUL, RSST-NIDS (confidence); UNAD+ 2026 (human check, post-hoc LIME); CADE 2021 (explains drift samples for analysts)",
     "No paper uses explanations to accept or reject pseudo-labels. Every existing gate uses confidence, model agreement or a human.",
     "Claim: explanation consistency as the admission rule for pseudo-labelled zero-day samples in continual NIDS. Must beat a confidence gate (PseudoFilter) as baseline. Motivation evidence: AOC-IDS 90% -> 75% collapse; GARDIAN on UGRansome (82.5% labels hurt fine-tuning).",
     "Strong: core contribution"),
    ("G2", GAPS["G2"], "PARTIALLY ADDRESSED",
     "DeepXplain 2026 (explanations in DRL reward / loss); C-MADF 2026 (explainability score gates actions); Drago et al. 2026 and CMC 2023 (SHAP-profile change triggers adaptation, no RL); GARDIAN (DRL retraining control, no XAI)",
     "Explanation signals inside DRL exist, but only for response actions; XAI-triggered adaptation exists, but without RL.",
     "Claim must be narrow: a DRL controller for continual retraining of a NIDS whose state includes explanation shift, with skip / rollback actions. Position explicitly against DeepXplain, C-MADF and Drago et al.",
     "Medium: good extension, not standalone"),
    ("G3", GAPS["G3"], "PARTIALLY ADDRESSED",
     "Drago et al. 2026 (explains drift impact and how decisions evolve); CADE 2021 (explains drift samples); CMC 2023 (SHAP drift detection); Explain to Not Forget, WIREs review (XAI against forgetting, outside IDS)",
     "Explaining drift in NIDS is done. Explaining forgetting caused by each continual update (replay / generated samples) in a NIDS is not.",
     "Use as an evaluation / analysis component (attribution diff before vs after each update), not as the main novelty.",
     "Weak as standalone"),
    ("G4", GAPS["G4"], "OPEN (narrow)",
     "None. ACORN-IDS, AOC-IDS (one-hot) and ThreatFormer-IDS (embeddings) choose an encoding without studying it; Alkasassbeh et al. label-encode",
     "No study of how categorical encoding changes zero-day detectability.",
     "Supporting experiment: our UGRansome result (label encoding AUC 0.49 vs one-hot 0.94).",
     "Supporting finding"),
]

# ---------------- write sheets ----------------
wb = load_workbook(XLSX)
for name in ("All Papers - Gap Test", "Gap Duplication Summary"):
    if name in wb.sheetnames:
        del wb[name]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
FILL = {HIGH: "F8CBAD", DUP: "FF7C80", PART: "FFE699", LOW: "E2EFDA"}


def header(ws, cols):
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"


ws = wb.create_sheet("All Papers - Gap Test")
cols = [("ID", 5), ("Title", 45), ("Year", 6), ("Venue", 26), ("Venue type", 20),
        ("Unsupervised", 8), ("DRL", 6), ("Continual / drift", 9), ("XAI", 6), ("Related part of your work", 32)]
for g in GAPS:
    cols += [(f"{g} verdict", 13), (f"{g} why", 40)]
cols += [("Extraction source", 16), ("Link", 30)]
header(ws, cols)
for r, p in enumerate(P, 2):
    row = [r - 1, p["title"], p["year"], p["venue"], p["vtype"], p["U"], p["D"], p["C"], p["X"], p["part"]]
    for g in ("g1", "g2", "g3", "g4"):
        row += [p[g][0], p[g][1] or "-"]
    row += [p["src"], p["link"] or "Not retrieved"]
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
        if isinstance(v, str) and v in FILL:
            cell.fill = PatternFill("solid", fgColor=FILL[v])
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(P) + 1}"

# legend under the table
lr = len(P) + 3
ws.cell(row=lr, column=2, value="Verdict legend").font = Font(bold=True)
for i, (v, txt) in enumerate([(DUP, "does exactly what the gap proposes"),
                              (HIGH, "same core idea in a different setting or with a non-XAI / non-RL substitute"),
                              (PART, "shares one key ingredient"), (LOW, "touches the topic only"),
                              ("No overlap", "not related to this gap")], 1):
    c = ws.cell(row=lr + i, column=2, value=v)
    if v in FILL:
        c.fill = PatternFill("solid", fgColor=FILL[v])
    ws.cell(row=lr + i, column=4, value=txt)
for i, (g, d) in enumerate(GAPS.items(), lr + 7):
    ws.cell(row=i, column=2, value=g).font = Font(bold=True)
    ws.cell(row=i, column=4, value=d)

ws2 = wb.create_sheet("Gap Duplication Summary")
header(ws2, [("Gap", 6), ("Definition", 45), ("Status", 18), ("Closest papers", 55), ("What exists vs what is missing", 50),
             ("How to claim novelty", 55), ("Strength", 18)])
counts = {g: {} for g in ("g1", "g2", "g3", "g4")}
for p in P:
    for g in counts:
        counts[g][p[g][0]] = counts[g].get(p[g][0], 0) + 1
for r, row in enumerate(SUMMARY, 2):
    for c, v in enumerate(row, 1):
        cell = ws2.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
    ws2.cell(row=r, column=3).fill = PatternFill("solid", fgColor="C6EFCE" if row[2].startswith("OPEN") else "FFE699")
r = len(SUMMARY) + 3
ws2.cell(row=r, column=1, value=f"Counts over all {len(P)} papers").font = Font(bold=True)
for i, g in enumerate(counts, 1):
    txt = ", ".join(f"{k}: {v}" for k, v in sorted(counts[g].items(), key=lambda kv: kv[0]))
    ws2.cell(row=r + i, column=1, value=g.upper())
    ws2.cell(row=r + i, column=2, value=txt)

wb.move_sheet("Gap Duplication Summary", offset=-(len(wb.sheetnames) - 2))
wb.move_sheet("All Papers - Gap Test", offset=-(len(wb.sheetnames) - 3))
wb.save(XLSX)
json.dump(dict(gaps=GAPS, papers=P, summary=SUMMARY), open(os.path.join(HERE, "gap_test.json"), "w"), indent=2, ensure_ascii=False)
print(len(P), "papers;", {g: counts[g] for g in counts})
print(wb.sheetnames)
